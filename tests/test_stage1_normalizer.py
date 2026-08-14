from __future__ import annotations

import json
from pathlib import Path
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace
from unittest.mock import patch

import stage1.normalizer as normalizer_module

from stage1.normalization_memory import NormalizationMemory, normalize_marker
from stage1.normalizer import (
    NormalizationValidationError,
    RevisionConflictError,
    SourceChangedError,
    confirm_normalization,
    generate_normalization,
    import_legacy_normalization,
    is_normalization_confirmed,
    load_normalization,
    preview_learning_candidates,
    reset_normalization,
    save_normalization,
)


class NormalizationMemoryTests(unittest.TestCase):
    def test_marker_normalization_and_manual_rule(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            db = Path(tmp) / "memory.sqlite3"
            with NormalizationMemory(db) as memory:
                self.assertEqual(normalize_marker(" Диктор-1: "), "диктор")
                created = memory.record_manual("Голос автора", "voiceover")
                self.assertEqual(created["status"], "active")
                resolved = memory.resolve_exact("голос автора:")
                self.assertIsNotNone(resolved)
                self.assertEqual(resolved["rule"]["target_type"], "voiceover")

    def test_llm_candidate_activates_after_distinct_confirmations(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            with NormalizationMemory(Path(tmp) / "memory.sqlite3") as memory:
                first = memory.record_llm_confirmation("Автор", "voiceover", "source-a")
                duplicate = memory.record_llm_confirmation("Автор", "voiceover", "source-a")
                second = memory.record_llm_confirmation("Автор", "voiceover", "source-b")
                duplicate_after_activation = memory.record_llm_confirmation(
                    "Автор", "voiceover", "source-b"
                )
                self.assertEqual(first["status"], "candidate")
                self.assertEqual(duplicate["confirmations"], 1)
                self.assertEqual(second["status"], "active")
                self.assertEqual(second["confirmations"], 2)
                self.assertEqual(duplicate_after_activation["status"], "active")

    def test_conflicting_manual_rules_disable_resolution(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            with NormalizationMemory(Path(tmp) / "memory.sqlite3") as memory:
                memory.record_manual("Текст", "voiceover")
                memory.record_manual("Текст", "sync")
                result = memory.resolve_exact("ТЕКСТ")
                self.assertTrue(result["conflict"])
                self.assertEqual({row["status"] for row in memory.list_rules()}, {"conflict"})

    def test_fuzzy_alias_is_suggestion_only(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            with NormalizationMemory(Path(tmp) / "memory.sqlite3") as memory:
                memory.record_manual("голос диктора", "voiceover")
                self.assertIsNone(memory.resolve_exact("голас диктора"))
                suggestion = memory.suggest("голас диктора")
                self.assertEqual(suggestion["rule"]["target_type"], "voiceover")

    def test_manual_rule_overrides_conflicting_llm_rule(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            with NormalizationMemory(Path(tmp) / "memory.sqlite3") as memory:
                memory.record_llm_confirmation("Автор", "sync", "source-a")
                memory.record_llm_confirmation("Автор", "sync", "source-b")
                memory.record_manual("Автор", "voiceover")

                resolved = memory.resolve_exact("автор")

                self.assertFalse(resolved["conflict"])
                self.assertEqual(resolved["rule"]["source"], "manual")
                self.assertEqual(resolved["rule"]["target_type"], "voiceover")

    def test_disabled_llm_rule_stays_disabled(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            with NormalizationMemory(Path(tmp) / "memory.sqlite3") as memory:
                rule = memory.record_llm_confirmation("Автор", "voiceover", "source-a")
                memory.update_rule(rule["id"], status="disabled")
                repeated = memory.record_llm_confirmation("Автор", "voiceover", "source-b")

                self.assertEqual(repeated["status"], "disabled")
                self.assertIsNone(memory.resolve_exact("Автор"))


class NormalizerTests(unittest.TestCase):
    def test_load_preserves_stage4_block_versioning_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "normalized.json"
            path.write_text(json.dumps({
                "source": "/tmp/scenario.docx",
                "source_hash": "abc",
                "blocks": [{
                    "uid": "stable-block",
                    "order": 1,
                    "original_text": "Изменённый текст",
                    "normalized_text": "Изменённый текст",
                    "versioning": {
                        "import_uid": "import-block",
                        "parent_block_uid": "stable-block",
                        "change_status": "edited",
                    },
                }],
            }), encoding="utf-8")

            loaded = load_normalization(path)

            self.assertEqual(loaded["blocks"][0]["versioning"], {
                "import_uid": "import-block",
                "parent_block_uid": "stable-block",
                "change_status": "edited",
            })

    def test_docx_soft_line_break_stays_inside_voiceover_block(self) -> None:
        from docx import Document
        from docx.enum.text import WD_BREAK

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.docx"
            document = Document()
            paragraph = document.add_paragraph("ДИКТОР: Первая строка")
            paragraph.add_run().add_break(WD_BREAK.LINE)
            paragraph.add_run("Вторая строка")
            document.save(source)

            result = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )

            self.assertEqual(len(result["blocks"]), 1)
            self.assertEqual(result["blocks"][0]["type"], "voiceover")
            self.assertEqual(result["blocks"][0]["normalized_text"], "Первая строка\nВторая строка")

    def test_docx_new_paragraph_continues_voiceover_without_repeating_marker(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.docx"
            document = Document()
            document.add_paragraph("ДИКТОР: Первый абзац")
            document.add_paragraph("Второй абзац того же закадра")
            document.add_paragraph("СНХ: Ответ героя")
            document.save(source)

            result = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3",
                use_llm=False,
            )

            self.assertEqual(
                [block["type"] for block in result["blocks"]],
                ["voiceover", "voiceover", "sync"],
            )
            self.assertEqual(
                result["blocks"][1]["text_for_tts"],
                "Второй абзац того же закадра",
            )

    def test_unlabelled_paragraph_is_voiceover_continuation_without_llm(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text("ДИКТОР: Первый абзац\nПродолжение закадра\n", encoding="utf-8")

            result = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3",
                use_llm=False,
            )
            self.assertEqual([block["type"] for block in result["blocks"]], ["voiceover", "voiceover"])
            self.assertEqual(result["blocks"][1]["continuation_of_type"], "voiceover")
            self.assertEqual(result["blocks"][1]["text_for_tts"], "Продолжение закадра")

    def test_standalone_voiceover_marker_is_non_tts_cue_for_next_paragraph(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text("З.к.\nСам закадровый текст.\n", encoding="utf-8")

            blocks = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"]

            self.assertEqual([block["type"] for block in blocks], ["technical", "voiceover"])
            self.assertEqual(blocks[0]["text_for_tts"], "")
            self.assertEqual(blocks[0]["notes"], "section_marker_for=voiceover")
            self.assertEqual(blocks[1]["text_for_tts"], "Сам закадровый текст.")

    def test_generate_preserves_every_paragraph_and_order(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            output = root / "normalized.json"
            source.write_text(
                "ДИКТОР: Первый закадр\nСНХ: Ответ эксперта\n(Показать карту)\nНеясный абзац\n",
                encoding="utf-8",
            )
            result = generate_normalization(
                source, output, memory_path=root / "memory.sqlite3", use_llm=False
            )
            self.assertEqual(len(result["blocks"]), 4)
            self.assertEqual([item["order"] for item in result["blocks"]], [1, 2, 3, 4])
            self.assertEqual(result["blocks"][0]["type"], "voiceover")
            self.assertEqual(result["blocks"][1]["type"], "sync")
            self.assertEqual(result["blocks"][2]["type"], "technical")
            self.assertEqual(result["blocks"][3]["type"], "unknown")
            self.assertEqual(result["blocks"][0]["original_text"], "ДИКТОР: Первый закадр")

    def test_regenerate_active_revision_preserves_identity_and_increments_revision(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            output = root / "normalized.json"
            source.write_text("ДИКТОР: Текст\n", encoding="utf-8")
            first = generate_normalization(
                source, output, memory_path=root / "memory.sqlite3", use_llm=False
            )
            first["scenario_revision_id"] = "rev_0003"
            first["parent_revision_id"] = "rev_0002"
            first["blocks"][0]["versioning"] = {"change_status": "unchanged"}
            output.write_text(json.dumps(first, ensure_ascii=False), encoding="utf-8")

            regenerated = generate_normalization(
                source, output, memory_path=root / "memory.sqlite3", use_llm=False
            )

            self.assertEqual(regenerated["revision"], 2)
            self.assertEqual(regenerated["scenario_revision_id"], "rev_0003")
            self.assertEqual(regenerated["parent_revision_id"], "rev_0002")
            self.assertEqual(regenerated["blocks"][0]["uid"], first["blocks"][0]["uid"])
            self.assertEqual(
                regenerated["blocks"][0]["versioning"],
                {"change_status": "unchanged"},
            )

    def test_regenerate_active_revision_rejects_another_source(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            first_source = root / "first.txt"
            second_source = root / "second.txt"
            output = root / "normalized.json"
            first_source.write_text("ДИКТОР: Первый\n", encoding="utf-8")
            second_source.write_text("ДИКТОР: Второй\n", encoding="utf-8")
            first = generate_normalization(
                first_source, output, memory_path=root / "memory.sqlite3", use_llm=False
            )
            first["scenario_revision_id"] = "rev_0001"
            output.write_text(json.dumps(first, ensure_ascii=False), encoding="utf-8")

            with self.assertRaises(SourceChangedError):
                generate_normalization(
                    second_source,
                    output,
                    memory_path=root / "memory.sqlite3",
                    use_llm=False,
                )

    def test_memory_applies_across_scenarios(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            db = root / "memory.sqlite3"
            with NormalizationMemory(db) as memory:
                memory.record_manual("Голос автора", "voiceover")
            source = root / "scenario.txt"
            source.write_text("ГОЛОС АВТОРА: Текст\n", encoding="utf-8")
            result = generate_normalization(source, root / "out.json", memory_path=db, use_llm=False)
            self.assertEqual(result["blocks"][0]["type"], "voiceover")
            self.assertEqual(result["blocks"][0]["source"], "memory")

    def test_text_only_edit_does_not_create_learning_candidate(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            normalized = root / "normalized.json"
            db = root / "memory.sqlite3"
            source.write_text("ДИКТОР: Исходный текст\n", encoding="utf-8")
            draft = generate_normalization(source, normalized, memory_path=db, use_llm=False)
            draft["blocks"][0]["normalized_text"] = "Исправленный текст"
            saved = save_normalization(normalized, draft, expected_revision=1)
            self.assertEqual(saved["blocks"][0]["source"], "builtin_rule")
            self.assertEqual(
                preview_learning_candidates(normalized, memory_path=db, expected_revision=2), []
            )

    def test_two_accepted_corrections_activate_learned_alias(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            db = root / "memory.sqlite3"
            for index in (1, 2):
                source = root / f"scenario-{index}.txt"
                normalized = root / f"normalized-{index}.json"
                source.write_text(f"ГОЛОС ХРОНИКИ: Текст {index}\n", encoding="utf-8")
                draft = generate_normalization(source, normalized, memory_path=db, use_llm=False)
                draft["blocks"][0]["type"] = "voiceover"
                draft["blocks"][0]["text_for_tts"] = f"Текст {index}"
                saved = save_normalization(normalized, draft, expected_revision=1)
                candidates = preview_learning_candidates(
                    normalized, memory_path=db, expected_revision=saved["revision"]
                )
                self.assertEqual(len(candidates), 1)
                confirm_normalization(
                    normalized, root / f"prepared-{index}.json", root / f"raw-{index}.txt",
                    memory_path=db, expected_revision=saved["revision"],
                    learning_decisions=[{
                        "candidate_id": candidates[0]["id"], "action": "remember",
                    }],
                )

            with NormalizationMemory(db) as memory:
                resolved = memory.resolve_exact("Голос хроники")
                self.assertIsNotNone(resolved)
                self.assertEqual(resolved["rule"]["source"], "correction")
                self.assertEqual(resolved["rule"]["confirmations"], 2)

    def test_never_decision_suppresses_future_learning_prompt(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            db = root / "memory.sqlite3"

            def corrected_draft(index: int):
                source = root / f"never-{index}.txt"
                normalized = root / f"never-{index}.json"
                source.write_text(f"ГОЛОС ЭПОХИ: Текст {index}\n", encoding="utf-8")
                draft = generate_normalization(source, normalized, memory_path=db, use_llm=False)
                draft["blocks"][0]["type"] = "voiceover"
                draft["blocks"][0]["text_for_tts"] = f"Текст {index}"
                saved = save_normalization(normalized, draft, expected_revision=1)
                return normalized, saved

            normalized, saved = corrected_draft(1)
            candidates = preview_learning_candidates(
                normalized, memory_path=db, expected_revision=saved["revision"]
            )
            confirm_normalization(
                normalized, root / "never-prepared.json", root / "never-raw.txt",
                memory_path=db, expected_revision=saved["revision"],
                learning_decisions=[{
                    "candidate_id": candidates[0]["id"], "action": "never",
                }],
            )
            second_path, second = corrected_draft(2)
            self.assertEqual(
                preview_learning_candidates(
                    second_path, memory_path=db, expected_revision=second["revision"]
                ),
                [],
            )

    def test_sync_metadata_is_normalized_from_known_marker(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "СНХ: Иван Иванов, эксперт, 01.02.03 Ответ героя\n",
                encoding="utf-8",
            )

            result = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )
            block = result["blocks"][0]

            self.assertEqual(block["type"], "sync")
            self.assertEqual(block["speaker"], "Иван Иванов")
            self.assertEqual(block["role"], "эксперт")
            self.assertEqual(block["timecode_start"], "01.02.03")
            self.assertEqual(block["normalized_text"], "Ответ героя")
            self.assertEqual(block["sync_fragments"], [{
                "text": "Ответ героя",
                "timecode_start": "01.02.03",
            }])

    def test_sync_speaker_block_can_contain_multiple_timecoded_fragments(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "ИРИНА КОТОВА, ИСТОРИК: 01:22:29. Шодэ приехал к нам из Одессы, где его "
                "семья жила с середины 19-ого века. Образование он получил во Франции "
                "художественное. 01:23:36. Был он человеком очень восприимчивым к разным "
                "стилям. 01:24:11. Он строил особняки многофасадные.\n",
                encoding="utf-8",
            )

            result = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )
            block = result["blocks"][0]

            self.assertEqual(block["type"], "sync")
            self.assertEqual(block["speaker"], "ИРИНА КОТОВА")
            self.assertEqual(block["role"], "ИСТОРИК")
            self.assertEqual(block["timecode_start"], "01:22:29")
            self.assertEqual(block["timecode_end"], "")
            self.assertEqual(block["sync_fragments"], [
                {
                    "text": (
                        "Шодэ приехал к нам из Одессы, где его семья жила с середины "
                        "19-ого века. Образование он получил во Франции художественное."
                    ),
                    "timecode_start": "01:22:29",
                },
                {
                    "text": "Был он человеком очень восприимчивым к разным стилям.",
                    "timecode_start": "01:23:36",
                },
                {
                    "text": "Он строил особняки многофасадные.",
                    "timecode_start": "01:24:11",
                },
            ])

    def test_sync_fragments_survive_docx_soft_line_breaks(self) -> None:
        from docx import Document
        from docx.enum.text import WD_BREAK

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.docx"
            document = Document()
            paragraph = document.add_paragraph(
                "ИРИНА КОТОВА, ИСТОРИК: 01:22:29. Первый фрагмент"
            )
            paragraph.add_run().add_break(WD_BREAK.LINE)
            paragraph.add_run("01:23:36. Второй фрагмент")
            paragraph.add_run().add_break(WD_BREAK.LINE)
            paragraph.add_run("01:24:11. Третий фрагмент")
            document.save(source)

            result = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )

            self.assertEqual(len(result["blocks"]), 1)
            self.assertEqual(
                [item["timecode_start"] for item in result["blocks"][0]["sync_fragments"]],
                ["01:22:29", "01:23:36", "01:24:11"],
            )

    def test_long_formal_speaker_role_is_sync_without_llm(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "ЕЛЕНА САДОВНИК, СТАРШИЙ НАУЧНЫЙ СОТРУДНИК МУЗЕЯ "
                "СИМБИРСКОГО КУПЕЧЕСТВА: 00:24:54. Образование он получил дома.\n",
                encoding="utf-8",
            )

            block = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"][0]

            self.assertEqual(block["type"], "sync")
            self.assertEqual(block["speaker"], "ЕЛЕНА САДОВНИК")
            self.assertEqual(block["timecode_start"], "00:24:54")

    def test_name_only_and_continuation_note_speaker_forms_are_sync(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "ЕЛЕНА САДОВНИК: 00:31:16. Семейная история.\n"
                "ИРИНА КОТОВА: (Продолжение синхрона) 01:25:03. Продолжение ответа.\n",
                encoding="utf-8",
            )

            blocks = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"]

            self.assertEqual([block["type"] for block in blocks], ["sync", "sync"])
            self.assertEqual([block["speaker"] for block in blocks], ["ЕЛЕНА САДОВНИК", "ИРИНА КОТОВА"])
            self.assertEqual([block["timecode_start"] for block in blocks], ["00:31:16", "01:25:03"])

    def test_known_voiceover_marker_precedes_generic_dash_label(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "ЗК Стратегия Лили — Она учится слышать собственный голос.\n",
                encoding="utf-8",
            )

            block = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"][0]

            self.assertEqual(block["type"], "voiceover")
            self.assertEqual(block["original_label"], "ЗК")
            self.assertEqual(
                block["normalized_text"],
                "Стратегия Лили — Она учится слышать собственный голос.",
            )
            self.assertEqual(block["text_for_tts"], block["normalized_text"])

    def test_dotted_and_glued_voiceover_markers_from_editorial_draft(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "З.к. Первый закадровый абзац.\n"
                "З.КВ 1915 году произошла встреча.\n"
                "З. Третий закадровый абзац.\n",
                encoding="utf-8",
            )

            blocks = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"]

            self.assertEqual([block["type"] for block in blocks], ["voiceover"] * 3)
            self.assertEqual(
                [block["normalized_text"] for block in blocks],
                [
                    "Первый закадровый абзац.",
                    "В 1915 году произошла встреча.",
                    "Третий закадровый абзац.",
                ],
            )

    def test_speaker_can_be_glued_to_timecode_or_use_spaced_three_part_timecode(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "СТОРОЖЕВА00:04:00 Первый ответ эксперта о футуристах.\n"
                "ИРБЕ. 00.28 45 Второй ответ эксперта о семье Брик.\n",
                encoding="utf-8",
            )

            blocks = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"]

            self.assertEqual([block["type"] for block in blocks], ["sync", "sync"])
            self.assertEqual([block["speaker"] for block in blocks], ["СТОРОЖЕВА", "ИРБЕ"])
            self.assertEqual(
                [block["timecode_start"] for block in blocks],
                ["00:04:00", "00:28:45"],
            )

    def test_long_all_caps_speaker_prefix_without_timecode_is_sync(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "КРАСНОВА Она объясняет происхождение галстука и подробно описывает экспонат музея.\n",
                encoding="utf-8",
            )

            block = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"][0]

            self.assertEqual(block["type"], "sync")
            self.assertEqual(block["speaker"], "КРАСНОВА")
            self.assertEqual(
                block["normalized_text"],
                "Она объясняет происхождение галстука и подробно описывает экспонат музея.",
            )

    def test_mixed_name_and_glued_period_speaker_prefixes_are_sync(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "Мария СМОЛЬНИКОВ Я благодарна Лиле за эту роль, потому что такой героини у меня не было.\n"
                "КРАСНОЩЕКОВА.Про Маяковского и Бриков всегда очень много сплетничали современники.\n",
                encoding="utf-8",
            )

            blocks = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"]

            self.assertEqual([block["type"] for block in blocks], ["sync", "sync"])
            self.assertEqual(
                [block["speaker"] for block in blocks],
                ["Мария СМОЛЬНИКОВ", "КРАСНОЩЕКОВА"],
            )

    def test_structural_heading_is_technical_not_voiceover_continuation(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "ЗК Основной текст.\nЧАСТЬ ВТОРАЯ. РЕПЕТИЦИЯ\nЗАСТАВКА МУЗА\n10.\n",
                encoding="utf-8",
            )

            blocks = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"]

            self.assertEqual(
                [block["type"] for block in blocks],
                ["voiceover", "technical", "technical", "technical"],
            )

    def test_single_word_speaker_can_precede_timecode_without_colon(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "Бартон 00:37 — 01:07 Я понял, что этот материал необходимо сохранить.\n",
                encoding="utf-8",
            )

            block = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"][0]

            self.assertEqual(block["type"], "sync")
            self.assertEqual(block["speaker"], "Бартон")
            self.assertEqual(block["timecode_start"], "00:37")
            self.assertEqual(block["timecode_end"], "01:07")
            self.assertEqual(
                block["normalized_text"],
                "Я понял, что этот материал необходимо сохранить.",
            )

    def test_speaker_role_can_precede_timecode_without_colon(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "Мария Смольникова, исполнительница роли Лили 40:45 — 41:00 "
                "Она всё время выбирает между долгом и собой.\n",
                encoding="utf-8",
            )

            block = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"][0]

            self.assertEqual(block["type"], "sync")
            self.assertEqual(block["speaker"], "Мария Смольникова")
            self.assertEqual(block["role"], "исполнительница роли Лили")
            self.assertEqual(block["timecode_start"], "40:45")
            self.assertEqual(block["timecode_end"], "41:00")

    def test_speaker_role_with_additional_commas_is_sync(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "Дана Курская, поэт, издатель: 17:38 — 18:05 "
                "В этой истории особенно важна интонация.\n",
                encoding="utf-8",
            )

            block = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"][0]

            self.assertEqual(block["type"], "sync")
            self.assertEqual(block["speaker"], "Дана Курская")
            self.assertEqual(block["role"], "поэт, издатель")
            self.assertEqual(block["timecode_start"], "17:38")
            self.assertEqual(block["timecode_end"], "18:05")

    def test_hyphenated_short_timecode_is_canonicalized(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "Шпагин 18-35 Это решение изменило дальнейшую работу.\n",
                encoding="utf-8",
            )

            block = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"][0]

            self.assertEqual(block["type"], "sync")
            self.assertEqual(block["speaker"], "Шпагин")
            self.assertEqual(block["timecode_start"], "18:35")
            self.assertEqual(block["sync_fragments"], [{
                "text": "Это решение изменило дальнейшую работу.",
                "timecode_start": "18:35",
            }])

    def test_prose_with_em_dash_is_not_split_into_a_label(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "Её главный выбор — услышать собственный голос.\n",
                encoding="utf-8",
            )

            block = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"][0]

            self.assertEqual(block["type"], "unknown")
            self.assertEqual(block["original_label"], "")
            self.assertEqual(
                block["normalized_text"],
                "Её главный выбор — услышать собственный голос.",
            )

    def test_clear_unlabelled_sync_continuations_inherit_speaker_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "Александра Ирбе, автор лекции-экскурсии: 19:03 — 19:50 Первый фрагмент.\n"
                "... я лично ничего в этом особенного не вижу.\n"
                "Смольникова 34:57 — 35:51 Начало мысли, которая может быть\n"
                "прелестной и одновременно сложной.\n"
                "Ирбе: 00:56 Первый фрагмент о доме.\n"
                "… некоторое время они ночевали на полу.\n",
                encoding="utf-8",
            )

            blocks = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"]

            self.assertEqual([blocks[index]["type"] for index in (1, 3, 5)], ["sync"] * 3)
            self.assertEqual(
                [blocks[index]["continuation_of_type"] for index in (1, 3, 5)],
                ["sync"] * 3,
            )
            self.assertEqual(blocks[1]["speaker"], "Александра Ирбе")
            self.assertEqual(blocks[1]["role"], "автор лекции-экскурсии")
            self.assertEqual(blocks[3]["speaker"], "Смольникова")
            self.assertEqual(blocks[5]["speaker"], "Ирбе")

    def test_capitalized_paragraph_after_sync_is_not_assumed_to_continue_it(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "Бартон 00:37 — 01:07 Ответ героя.\n"
                "Самостоятельный редакционный абзац без маркера.\n",
                encoding="utf-8",
            )

            blocks = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"]

            self.assertEqual(blocks[0]["type"], "sync")
            self.assertEqual(blocks[1]["type"], "unknown")
            self.assertEqual(blocks[1]["continuation_of_type"], "")

    def test_connective_paragraph_after_sync_inherits_speaker_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "Курская: 04-59 — 5:39 Начало ответа.\n"
                "И литературная тусовка слушается, для неё это важно.\n",
                encoding="utf-8",
            )

            blocks = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"]

            self.assertEqual([block["type"] for block in blocks], ["sync", "sync"])
            self.assertEqual(blocks[1]["continuation_of_type"], "sync")
            self.assertEqual(blocks[1]["speaker"], "Курская")

    def test_bare_snh_marker_inherits_the_previous_named_sync(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "Бартон 00:37 — 01:07 Первый выбранный фрагмент.\n"
                "СНХ: Фраза, вырезанная из другого места интервью.\n"
                "СНХ: Фраза, которую сценарист переставил раньше исходника.\n",
                encoding="utf-8",
            )

            blocks = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"]

            self.assertEqual([block["type"] for block in blocks], ["sync"] * 3)
            self.assertEqual([block["speaker"] for block in blocks], ["Бартон"] * 3)
            self.assertEqual(
                [block["continuation_of_type"] for block in blocks],
                ["", "sync", "sync"],
            )

    def test_parenthetical_editorial_prefix_before_timecode_is_sync(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            lines = [
                "(у дома в Спасопесковском): 01:04:20 - Здесь был кооператив.",
                "(или более эмоционально Шпагин: 30:37 — я думаю, Лиля написала письмо.",
            ]
            source.write_text("\n".join(lines) + "\n", encoding="utf-8")

            blocks = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"]

            self.assertEqual([block["type"] for block in blocks], ["sync", "sync"])
            self.assertEqual([block["timecode_start"] for block in blocks], ["01:04:20", "30:37"])
            self.assertEqual([block["original_label"] for block in blocks], ["", ""])
            self.assertEqual([block["normalized_text"] for block in blocks], lines)

    def test_voiceover_prose_chains_keep_full_text(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            lines = [
                "ЗК Маяковский собирался жить в новой квартире с Полонской.",
                "Но Нора была уже замужем — за Михаилом Яншиным.",
                "В феврале Брики — которые давно не ездили — вдруг уехали в Европу.",
                "ЗК Лиля вытребовала себе половину авторских отчислений.",
                "Видимо, она считала, что имеет право — как вдова?",
                "Но сначала большими деньгами не пахло: Маяковского издавали крайне мало.",
            ]
            source.write_text("\n".join(lines) + "\n", encoding="utf-8")

            blocks = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"]

            self.assertEqual([block["type"] for block in blocks], ["voiceover"] * 6)
            for index in (1, 2, 4, 5):
                self.assertEqual(blocks[index]["original_label"], "")
                self.assertEqual(blocks[index]["normalized_text"], lines[index])
                self.assertEqual(blocks[index]["text_for_tts"], lines[index])

    def test_unprefixed_shot_plan_is_technical_not_voiceover_continuation(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "ЗАКАДР: Основной текст.\nПланы дома Шатрова\nПортреты архитектора Шодэ\n",
                encoding="utf-8",
            )

            blocks = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"]

            self.assertEqual([block["type"] for block in blocks], [
                "voiceover", "technical", "technical",
            ])

    def test_plural_titles_marker_is_technical(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text("ТИТРЫ\n", encoding="utf-8")

            block = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )["blocks"][0]

            self.assertEqual(block["type"], "technical")

    def test_sync_fragment_explicit_range_has_optional_end_timecode(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "СНХ: Иван Иванов, эксперт, 01:02:03–01:02:08 Ответ героя\n",
                encoding="utf-8",
            )

            result = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )
            block = result["blocks"][0]

            self.assertEqual(block["timecode_start"], "01:02:03")
            self.assertEqual(block["timecode_end"], "01:02:08")
            self.assertEqual(block["sync_fragments"], [{
                "text": "Ответ героя",
                "timecode_start": "01:02:03",
                "timecode_end": "01:02:08",
            }])

    def test_llm_classifies_only_unresolved_without_drops(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text("ДИКТОР: Готово\nАвтор: Новый текст\n", encoding="utf-8")

            def classifier(items, _hint):
                self.assertEqual(len(items), 1)
                return [{
                    "uid": items[0]["uid"], "type": "voiceover", "confidence": .8,
                    "label_is_marker": True,
                }]

            result = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3",
                classifier=classifier,
            )
            self.assertEqual(len(result["blocks"]), 2)
            self.assertEqual(result["blocks"][1]["source"], "llm")
            self.assertEqual(result["blocks"][1]["text_for_tts"], "Новый текст")

    def test_prose_prefix_is_not_stripped_without_marker_confirmation(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text("Москва: сегодня прошёл форум\n", encoding="utf-8")

            def classifier(items, _hint):
                return [{
                    "uid": items[0]["uid"], "type": "voiceover", "confidence": .9,
                    "label_is_marker": False,
                }]

            result = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3",
                classifier=classifier,
            )

            self.assertEqual(result["blocks"][0]["normalized_text"], "Москва: сегодня прошёл форум")
            self.assertEqual(result["blocks"][0]["text_for_tts"], "Москва: сегодня прошёл форум")
            self.assertEqual(result["blocks"][0]["original_label"], "")

    def test_manual_type_change_does_not_learn_or_strip_a_prose_prefix(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            normalized = root / "normalized.json"
            prepared = root / "prepared.json"
            raw = root / "raw.txt"
            memory_path = root / "memory.sqlite3"
            source.write_text("Москва: сегодня прошёл форум\n", encoding="utf-8")

            draft = generate_normalization(
                source, normalized, memory_path=memory_path, use_llm=False
            )
            draft["blocks"][0]["type"] = "voiceover"
            draft["blocks"][0]["text_for_tts"] = "Москва: сегодня прошёл форум"
            saved = save_normalization(normalized, draft, expected_revision=1)
            confirm_normalization(
                normalized,
                prepared,
                raw,
                memory_path=memory_path,
                expected_revision=saved["revision"],
            )

            self.assertIn("Москва: сегодня прошёл форум", raw.read_text(encoding="utf-8"))
            with NormalizationMemory(memory_path) as memory:
                self.assertIsNone(memory.resolve_exact("Москва"))

    def test_timecode_prefix_is_not_split_as_label(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text("00:01:02 Иван говорит\n", encoding="utf-8")
            result = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )
            self.assertEqual(result["blocks"][0]["original_label"], "")
            self.assertEqual(result["blocks"][0]["normalized_text"], "00:01:02 Иван говорит")

    def test_prose_leading_date_is_not_classified_as_sync(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text("23.05 состоялась премьера\n", encoding="utf-8")

            result = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )

            block = result["blocks"][0]
            self.assertEqual(block["type"], "unknown")
            self.assertNotEqual(block["confidence"], 0.9)

    def test_low_confidence_llm_result_requires_manual_resolution(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text("Авторский блок\n", encoding="utf-8")

            def classifier(items, _hint):
                return [{"uid": items[0]["uid"], "type": "voiceover", "confidence": .1}]

            result = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3",
                classifier=classifier,
            )
            self.assertEqual(result["blocks"][0]["type"], "unknown")
            self.assertTrue(result["blocks"][0]["warnings"])

    def test_large_scenario_is_classified_in_batches(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text("\n".join(["А" * 7000, "Б" * 7000, "В" * 7000]), encoding="utf-8")
            calls: list[int] = []

            def classifier(items, _hint):
                calls.append(len(items))
                return [
                    {"uid": item["uid"], "type": "other", "confidence": 0.7}
                    for item in items
                ]

            result = generate_normalization(
                source,
                root / "out.json",
                memory_path=root / "memory.sqlite3",
                classifier=classifier,
            )

            self.assertGreater(len(calls), 1)
            self.assertEqual(len(result["blocks"]), 3)
            self.assertEqual([item["type"] for item in result["blocks"]], ["other"] * 3)

    def test_default_llm_classifier_handles_multiple_compact_batches(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text(
                "\n".join(f"Абзац {index} " + "текст " * 90 for index in range(80)),
                encoding="utf-8",
            )

            def classifier(items, _hint):
                self.assertTrue(items)
                self.assertNotIn("normalized_text", items[0])
                return [
                    {"uid": item["uid"], "type": "other", "confidence": 0.9}
                    for item in items
                ]

            with patch.object(
                normalizer_module, "_openrouter_classifier", side_effect=classifier
            ) as mocked:
                result = generate_normalization(
                    source, root / "out.json", memory_path=root / "memory.sqlite3"
                )

            self.assertGreater(mocked.call_count, 1)
            self.assertEqual(len(result["blocks"]), 80)
            self.assertTrue(all(item["type"] == "other" for item in result["blocks"]))

    def test_llm_batch_cache_resumes_after_interrupted_batch(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            output = root / "normalized.json"
            memory_path = root / "memory.sqlite3"
            source.write_text(
                "\n".join(["А" * 7000, "Б" * 7000, "В" * 7000]),
                encoding="utf-8",
            )

            class FlakyClassifier:
                def __init__(self):
                    self.calls = 0
                    self.failed = False

                def __call__(self, items, _hint):
                    self.calls += 1
                    if self.calls == 2 and not self.failed:
                        self.failed = True
                        raise RuntimeError("provider interrupted")
                    return [
                        {"uid": item["uid"], "type": "other", "confidence": 0.9}
                        for item in items
                    ]

            classifier = FlakyClassifier()
            first = generate_normalization(
                source,
                output,
                memory_path=memory_path,
                classifier=classifier,
            )
            cache_path = root / "work" / "normalization_llm_cache.json"
            cache_after_first = json.loads(cache_path.read_text(encoding="utf-8"))

            self.assertEqual(classifier.calls, 3)
            self.assertEqual(len(cache_after_first["entries"]), 2)
            self.assertEqual(
                [block["type"] for block in first["blocks"]],
                ["other", "unknown", "other"],
            )

            resumed = generate_normalization(
                source,
                output,
                memory_path=memory_path,
                classifier=classifier,
            )
            cache_after_resume = json.loads(cache_path.read_text(encoding="utf-8"))

            self.assertEqual(classifier.calls, 4)
            self.assertEqual(len(cache_after_resume["entries"]), 3)
            self.assertTrue(all(block["type"] == "other" for block in resumed["blocks"]))
            self.assertTrue(all(block["source"] == "llm" for block in resumed["blocks"]))
            self.assertEqual(
                list(cache_path.parent.glob(".normalization_llm_cache.json.*")),
                [],
            )

    def test_llm_cache_does_not_create_nested_work_directory(self) -> None:
        output = Path("projects/demo/work/scenario_normalized.json")
        self.assertEqual(
            normalizer_module._llm_cache_path(output),
            Path("projects/demo/work/normalization_llm_cache.json"),
        )

    def test_llm_cache_invalidates_for_text_hint_and_memory_rules(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            output = root / "normalized.json"
            memory_path = root / "memory.sqlite3"
            calls = 0

            def classifier(items, _hint):
                nonlocal calls
                calls += 1
                return [
                    {"uid": item["uid"], "type": "other", "confidence": 0.9}
                    for item in items
                ]

            source.write_text("Авторский абзац", encoding="utf-8")
            generate_normalization(
                source, output, memory_path=memory_path, classifier=classifier
            )
            generate_normalization(
                source, output, memory_path=memory_path, classifier=classifier
            )
            self.assertEqual(calls, 1)

            source.write_text("Изменённый авторский абзац", encoding="utf-8")
            generate_normalization(
                source, output, memory_path=memory_path, classifier=classifier
            )
            self.assertEqual(calls, 2)

            generate_normalization(
                source,
                output,
                memory_path=memory_path,
                classifier=classifier,
                hint="Это интервью",
            )
            self.assertEqual(calls, 3)

            with NormalizationMemory(memory_path) as memory:
                memory.record_manual("Несвязанный маркер", "technical")
            generate_normalization(
                source,
                output,
                memory_path=memory_path,
                classifier=classifier,
                hint="Это интервью",
            )
            self.assertEqual(calls, 4)

    def test_default_llm_cache_invalidates_for_model_and_prompt(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            output = root / "normalized.json"
            memory_path = root / "memory.sqlite3"
            source.write_text("Авторский абзац", encoding="utf-8")
            current_model = {"value": "model-a"}

            def settings():
                return SimpleNamespace(
                    openrouter_model=current_model["value"],
                    openrouter_base_url="https://example.invalid/api/v1",
                )

            def classifier(items, _hint):
                return [
                    {"uid": item["uid"], "type": "other", "confidence": 0.9}
                    for item in items
                ]

            with (
                patch.object(normalizer_module, "get_settings", side_effect=settings),
                patch.object(
                    normalizer_module,
                    "_openrouter_classifier",
                    side_effect=classifier,
                ) as mocked,
            ):
                generate_normalization(source, output, memory_path=memory_path)
                generate_normalization(source, output, memory_path=memory_path)
                self.assertEqual(mocked.call_count, 1)

                current_model["value"] = "model-b"
                generate_normalization(source, output, memory_path=memory_path)
                self.assertEqual(mocked.call_count, 2)

                with patch.object(
                    normalizer_module,
                    "_OPENROUTER_INSTRUCTIONS",
                    normalizer_module._OPENROUTER_INSTRUCTIONS + "Prompt revision. ",
                ):
                    generate_normalization(source, output, memory_path=memory_path)
                self.assertEqual(mocked.call_count, 3)

    def test_save_revision_and_confirm_projections(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            normalized = root / "normalized.json"
            prepared = root / "prepared.json"
            raw = root / "raw.txt"
            db = root / "memory.sqlite3"
            source.write_text("Необычный голос: Текст\nСНХ: Ответ\n", encoding="utf-8")
            draft = generate_normalization(source, normalized, memory_path=db, use_llm=False)
            draft["blocks"][0]["type"] = "voiceover"
            draft["blocks"][0]["text_for_tts"] = "Текст"
            saved = save_normalization(normalized, draft, expected_revision=1)
            self.assertEqual(saved["revision"], 2)
            self.assertEqual(saved["blocks"][0]["source"], "manual")
            with self.assertRaises(RevisionConflictError):
                save_normalization(normalized, saved, expected_revision=1)

            candidates = preview_learning_candidates(
                normalized, memory_path=db, expected_revision=2
            )
            self.assertEqual(len(candidates), 1)
            confirmed = confirm_normalization(
                normalized, prepared, raw, memory_path=db, expected_revision=2,
                learning_decisions=[{
                    "candidate_id": candidates[0]["id"], "action": "remember",
                }],
            )
            self.assertEqual(confirmed["status"], "confirmed")
            self.assertTrue(is_normalization_confirmed(normalized))
            payload = json.loads(prepared.read_text(encoding="utf-8"))
            self.assertEqual(payload["items"][0]["id"], "zk1")
            self.assertEqual(payload["normalization_revision"], confirmed["revision"])
            self.assertIn("ЗК1: Текст", raw.read_text(encoding="utf-8"))
            with NormalizationMemory(db) as memory:
                rules = [r for r in memory.list_rules() if r["source"] == "correction"]
                self.assertEqual(len(rules), 1)
                self.assertEqual(rules[0]["target_type"], "voiceover")
                self.assertEqual(rules[0]["status"], "candidate")

    def test_reset_clears_blocks_and_is_revision_safe(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            normalized = root / "normalized.json"
            source.write_text("ДИКТОР: Исходный текст\nСНХ: Исходный ответ\n", encoding="utf-8")
            draft = generate_normalization(
                source, normalized, memory_path=root / "memory.sqlite3", use_llm=False
            )
            draft["blocks"].reverse()
            for index, block in enumerate(draft["blocks"], start=1):
                block["order"] = index
                block["normalized_text"] = "Ручная правка"
            saved = save_normalization(normalized, draft, expected_revision=1)

            reset = reset_normalization(
                normalized, expected_revision=saved["revision"], source_root=root
            )

            self.assertEqual(reset["revision"], saved["revision"] + 1)
            self.assertEqual(reset["status"], "draft")
            self.assertEqual(reset["hint"], "")
            self.assertEqual(reset["blocks"], [])
            with self.assertRaises(RevisionConflictError):
                reset_normalization(normalized, expected_revision=saved["revision"], source_root=root)

    def test_unknown_blocks_prevent_confirmation(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text("Непонятный абзац\n", encoding="utf-8")
            normalized = root / "normalized.json"
            generate_normalization(source, normalized, memory_path=root / "db.sqlite3", use_llm=False)
            with self.assertRaises(NormalizationValidationError):
                confirm_normalization(
                    normalized, root / "prepared.json", root / "raw.txt",
                    memory_path=root / "db.sqlite3",
                )
            self.assertFalse((root / "prepared.json").exists())

    def test_failed_publication_does_not_update_memory(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            normalized = root / "normalized.json"
            prepared = root / "prepared.json"
            raw = root / "raw.txt"
            db = root / "memory.sqlite3"
            source.write_text("НЕОБЫЧНЫЙ ГОЛОС: Текст\n", encoding="utf-8")
            draft = generate_normalization(source, normalized, memory_path=db, use_llm=False)
            draft["blocks"][0]["type"] = "voiceover"
            draft["blocks"][0]["text_for_tts"] = "Текст"
            saved = save_normalization(normalized, draft, expected_revision=1)
            candidates = preview_learning_candidates(
                normalized, memory_path=db, expected_revision=saved["revision"]
            )
            original_write = normalizer_module._atomic_write_text

            def fail_raw(path, content):
                if Path(path) == raw:
                    raise OSError("simulated raw projection failure")
                return original_write(path, content)

            with patch("stage1.normalizer._atomic_write_text", side_effect=fail_raw):
                with self.assertRaises(OSError):
                    confirm_normalization(
                        normalized, prepared, raw, memory_path=db, expected_revision=2,
                        learning_decisions=[{
                            "candidate_id": candidates[0]["id"], "action": "remember",
                        }],
                    )

            with NormalizationMemory(db) as memory:
                self.assertIsNone(memory.resolve_exact("НЕОБЫЧНЫЙ ГОЛОС"))

    def test_legacy_import_preserves_voiceover_sync_host_and_technical(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            source.write_text("source", encoding="utf-8")
            prepared = root / "prepared.json"
            prepared.write_text(json.dumps({
                "source": str(source),
                "items": [{"text_original": "Закадр", "text_for_tts": "Закадр"}],
            }), encoding="utf-8")
            raw = root / "raw.txt"
            raw.write_text(
                "ЗК1: Закадр\n\n"
                "СНХ Иван Иванов, эксперт, 01:02:03 Ответ\n\n"
                "ВЕДУЩИЙ: Реплика\n\nОТБИВКА\n",
                encoding="utf-8",
            )

            result = import_legacy_normalization(prepared, raw, root / "normalized.json", source_path=source)

            self.assertEqual(
                [block["type"] for block in result["blocks"]],
                ["voiceover", "sync", "host", "technical"],
            )

    def test_legacy_import_without_original_source_can_be_confirmed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            prepared = root / "prepared.json"
            raw = root / "raw.txt"
            normalized = root / "normalized.json"
            prepared.write_text(json.dumps({
                "source": str(root / "missing.docx"),
                "items": [{"text_original": "Текст", "text_for_tts": "Текст"}],
            }), encoding="utf-8")
            raw.write_text("ЗК1: Текст\n", encoding="utf-8")
            imported = import_legacy_normalization(prepared, raw, normalized)

            confirmed = confirm_normalization(
                normalized,
                root / "confirmed-prepared.json",
                root / "confirmed-raw.txt",
                memory_path=root / "memory.sqlite3",
                expected_revision=imported["revision"],
            )

            self.assertEqual(confirmed["status"], "confirmed")

    def test_parallel_saves_with_same_revision_do_not_both_succeed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.txt"
            normalized = root / "normalized.json"
            source.write_text("ДИКТОР: Текст\n", encoding="utf-8")
            draft = generate_normalization(
                source, normalized, memory_path=root / "memory.sqlite3", use_llm=False
            )

            def save_once(suffix):
                payload = json.loads(json.dumps(draft))
                payload["blocks"][0]["normalized_text"] += suffix
                try:
                    save_normalization(normalized, payload, expected_revision=1)
                    return "saved"
                except RevisionConflictError:
                    return "conflict"

            with ThreadPoolExecutor(max_workers=2) as pool:
                outcomes = list(pool.map(save_once, [" A", " B"]))

            self.assertEqual(sorted(outcomes), ["conflict", "saved"])

    def test_docx_table_cells_are_included_in_order(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scenario.docx"
            document = Document()
            document.add_paragraph("ДИКТОР: До таблицы")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "СНХ: Ответ"
            table.cell(0, 1).text = "ОТБИВКА"
            document.add_paragraph("ДИКТОР: После таблицы")
            document.save(source)

            result = generate_normalization(
                source, root / "out.json", memory_path=root / "memory.sqlite3", use_llm=False
            )

            self.assertEqual(len(result["blocks"]), 4)
            self.assertEqual(
                [block["type"] for block in result["blocks"]],
                ["voiceover", "sync", "technical", "voiceover"],
            )


if __name__ == "__main__":
    unittest.main()
