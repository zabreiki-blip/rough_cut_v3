from __future__ import annotations

import json
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import threading
import unittest
from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from stage1.analyzer import (
    _cleanup_blocks,
    _extract_manual_narrator_blocks,
    _merge_blocks_to_limit,
    analyze_script,
)
from stage1.models import NarrationBlock


class Stage1AnalyzerTests(unittest.TestCase):
    def test_extract_manual_blocks_supports_decimal_and_letter_suffixes(self) -> None:
        scenario_text = """
ЗК4.1 За какой информацией из Арктики идёт охота сегодня?
Почему даже климатические исследования региона засекречены?
ВЕДУЩИЙ 1: Подводка в кадре.
ЗК38А
Чету Аугустенборгов выслали из СССР.
СНХ Юрий КНУТОВ, военный эксперт: 00:06:10 Реплика.
""".strip()

        blocks = _extract_manual_narrator_blocks(scenario_text)

        self.assertEqual(
            [block.text_for_tts for block in blocks],
            [
                "За какой информацией из Арктики идёт охота сегодня?\n"
                "Почему даже климатические исследования региона засекречены?",
                "Чету Аугустенборгов выслали из СССР.",
            ],
        )

    def test_cleanup_removes_markers_speakers_and_technical_lines(self) -> None:
        raw_text = "\n".join(
            [
                "ЗК25",
                "Нужный закадровый текст.",
                "СНХ Юрий КНУТОВ, военный эксперт: 00:06:10 Реплика героя.",
                "ЛАЙВ",
                "ВЕДУЩИЙ 1: Реквизит: показать фото.",
            ]
        )
        blocks = [NarrationBlock(id="", text_original=raw_text, text_for_tts=raw_text)]

        cleaned = _cleanup_blocks(blocks)

        self.assertEqual(len(cleaned), 1)
        self.assertEqual(cleaned[0].text_for_tts, "Нужный закадровый текст.")

    def test_highlighted_docx_uses_embedded_markers_to_split_blocks(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            tmp = Path(tmp_dir)
            input_path = tmp / "scenario.docx"
            output_path = tmp / "scenario_prepared.json"
            raw_output_path = tmp / "scenario_raw.txt"

            doc = Document()
            for paragraph_text in [
                "ЗК1",
                "Первый закадровый блок достаточно длинный для отдельного TTS.",
                "СНХ Юрий КНУТОВ, военный эксперт:",
                "00:06:10 Реплика героя, это не закадр.",
                "ЗК2 Второй закадровый блок тоже достаточно длинный и самостоятельный.",
            ]:
                paragraph = doc.add_paragraph()
                run = paragraph.add_run(paragraph_text)
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            doc.save(input_path)

            result = analyze_script(
                input_path=input_path,
                output_path=output_path,
                raw_output_path=raw_output_path,
                use_llm=False,
            )
            payload = json.loads(output_path.read_text(encoding="utf-8"))

            self.assertEqual(result["blocks_count"], 2)
            self.assertEqual(payload["extraction_method"], "manual_markers")
            self.assertEqual(
                [item["text_for_tts"] for item in payload["items"]],
                [
                    "Первый закадровый блок достаточно длинный для отдельного TTS.",
                    "Второй закадровый блок тоже достаточно длинный и самостоятельный.",
                ],
            )

    def test_merge_blocks_to_limit_prefers_short_adjacent_pairs(self) -> None:
        blocks = [
            NarrationBlock(id="", text_original="A" * 180, text_for_tts="A" * 180, notes="long_1"),
            NarrationBlock(id="", text_original="b" * 12, text_for_tts="b" * 12, notes="short_1"),
            NarrationBlock(id="", text_original="c" * 13, text_for_tts="c" * 13, notes="short_2"),
            NarrationBlock(id="", text_original="D" * 170, text_for_tts="D" * 170, notes="long_2"),
            NarrationBlock(id="", text_original="e" * 11, text_for_tts="e" * 11, notes="short_3"),
            NarrationBlock(id="", text_original="f" * 14, text_for_tts="f" * 14, notes="short_4"),
        ]

        merged = _merge_blocks_to_limit(blocks, max_blocks=4)

        self.assertEqual(len(merged), 4)
        self.assertEqual(merged[0].text_for_tts, "A" * 180)
        self.assertEqual(merged[1].text_for_tts, f"{'b' * 12} {'c' * 13}")
        self.assertEqual(merged[2].text_for_tts, "D" * 170)
        self.assertEqual(merged[3].text_for_tts, f"{'e' * 11} {'f' * 14}")
        self.assertEqual(merged[1].notes, "short_1 | short_2")
        self.assertEqual(merged[3].notes, "short_3 | short_4")

    def test_llm_analysis_writes_prepared_and_raw_outputs(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            tmp = Path(tmp_dir)
            input_path = tmp / "scenario.txt"
            output_path = tmp / "scenario_prepared.json"
            raw_output_path = tmp / "scenario_raw.txt"
            narrator_text = "Авторский текст для диктора достаточно длинный для теста."
            input_path.write_text(
                f"{narrator_text}\n\nСНХ Иван Иванов: реплика героя.\n",
                encoding="utf-8",
            )

            llm_blocks = [
                NarrationBlock(
                    id="zk1",
                    text_original=narrator_text,
                    text_for_tts=narrator_text,
                )
            ]
            with (
                patch.dict(os.environ, {"OPENROUTER_API_KEY": "test-key"}),
                patch("stage1.analyzer.analyze_with_openrouter", return_value=llm_blocks),
            ):
                result = analyze_script(
                    input_path=input_path,
                    output_path=output_path,
                    raw_output_path=raw_output_path,
                    use_llm=True,
                )

            payload = json.loads(output_path.read_text(encoding="utf-8"))
            self.assertTrue(result["used_llm"])
            self.assertEqual(result["extraction_method"], "llm")
            self.assertEqual(payload["items"][0]["text_for_tts"], narrator_text)
            self.assertIn(f"ЗК1: {narrator_text}", raw_output_path.read_text(encoding="utf-8"))

    def test_llm_chunk_failure_falls_back_and_reports_error(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            tmp = Path(tmp_dir)
            input_path = tmp / "scenario.txt"
            output_path = tmp / "scenario_prepared.json"
            raw_output_path = tmp / "scenario_raw.txt"
            input_path.write_text(
                "Первый авторский абзац достаточно длинный для fallback-проверки.",
                encoding="utf-8",
            )

            with (
                patch.dict(os.environ, {"OPENROUTER_API_KEY": "test-key"}),
                patch(
                    "stage1.analyzer.analyze_with_openrouter",
                    side_effect=RuntimeError("temporary LLM error"),
                ),
            ):
                result = analyze_script(
                    input_path=input_path,
                    output_path=output_path,
                    raw_output_path=raw_output_path,
                    use_llm=True,
                )

            payload = json.loads(output_path.read_text(encoding="utf-8"))
            self.assertFalse(result["used_llm"])
            self.assertEqual(result["extraction_method"], "fallback")
            self.assertEqual(result["llm_errors_count"], 1)
            self.assertIn("temporary LLM error", payload["llm_errors"][0])
            self.assertEqual(len(payload["items"]), 1)

    def test_llm_chunks_run_concurrently_but_keep_source_order(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            tmp = Path(tmp_dir)
            input_path = tmp / "scenario.txt"
            output_path = tmp / "scenario_prepared.json"
            raw_output_path = tmp / "scenario_raw.txt"
            input_path.write_text("source text", encoding="utf-8")
            chunks = ["first chunk", "second chunk"]
            both_started = threading.Event()
            lock = threading.Lock()
            started = 0

            def analyze_chunk(_settings, chunk):
                nonlocal started
                with lock:
                    started += 1
                    if started == 2:
                        both_started.set()
                if not both_started.wait(timeout=1):
                    raise RuntimeError("chunks were processed serially")
                text = "A" * 60 if chunk == "first chunk" else "B" * 60
                return [NarrationBlock(id="", text_original=text, text_for_tts=text)]

            with (
                patch.dict(
                    os.environ, {"OPENROUTER_API_KEY": "test-key", "ANALYZER_LLM_WORKERS": "2"}
                ),
                patch("stage1.analyzer._chunk_for_llm", return_value=chunks),
                patch("stage1.analyzer.analyze_with_openrouter", side_effect=analyze_chunk),
            ):
                result = analyze_script(
                    input_path=input_path,
                    output_path=output_path,
                    raw_output_path=raw_output_path,
                    use_llm=True,
                )

            payload = json.loads(output_path.read_text(encoding="utf-8"))
            self.assertTrue(result["used_llm"])
            self.assertEqual(
                [item["text_for_tts"] for item in payload["items"]],
                ["A" * 60, "B" * 60],
            )


if __name__ == "__main__":
    unittest.main()
