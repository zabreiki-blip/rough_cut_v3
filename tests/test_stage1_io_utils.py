from __future__ import annotations

from pathlib import Path
import subprocess
from subprocess import CompletedProcess
from types import SimpleNamespace
from unittest.mock import patch

import pytest

import stage1.io_utils as stage1_io_utils
from stage1.io_utils import (
    _paragraph_is_marked,
    _run_is_marked,
    _xml_has_marking,
    extract_marked_blocks_from_docx,
    extract_paragraphs_from_docx,
    extract_text,
    extract_text_from_doc,
)


def test_doc_extraction_prefers_antiword_without_visual_line_wrapping(tmp_path: Path) -> None:
    source = tmp_path / "scenario.doc"
    source.write_bytes(b"legacy-word-placeholder")
    extracted = (
        "ЗАКАДР: Один длинный логический абзац без искусственных переносов.\n"
        "ИРИНА КОТОВА, ИСТОРИК: 01:22:29. Реплика.\n"
    )

    with patch(
        "stage1.io_utils.subprocess.run",
        return_value=CompletedProcess([], 0, stdout=extracted, stderr=""),
    ) as run:
        result = extract_text_from_doc(source)

    assert run.call_args.args[0] == ["antiword", "-w", "0", str(source)]
    assert result.splitlines() == extracted.strip().splitlines()


# ---------------------------------------------------------------------------
# _run_is_marked / _paragraph_is_marked / _xml_has_marking
# ---------------------------------------------------------------------------


def test_xml_has_marking_returns_false_for_none_element() -> None:
    assert _xml_has_marking(None) is False


def test_run_is_marked_false_without_font_or_xml() -> None:
    run = SimpleNamespace(font=None, _element=None)
    assert _run_is_marked(run) is False


def test_run_is_marked_true_via_font_highlight_color() -> None:
    font = SimpleNamespace(highlight_color=object())
    run = SimpleNamespace(font=font, _element=None)
    assert _run_is_marked(run) is True


def test_paragraph_is_marked_false_without_xml_element() -> None:
    paragraph = SimpleNamespace(_p=None)
    assert _paragraph_is_marked(paragraph) is False


def test_xml_has_marking_detects_highlight_tag() -> None:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    element = parse_xml(f'<w:rPr {nsdecls("w")}><w:highlight w:val="yellow"/></w:rPr>')
    assert _xml_has_marking(element) is True


def test_xml_has_marking_ignores_highlight_none() -> None:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    element = parse_xml(f'<w:rPr {nsdecls("w")}><w:highlight w:val="none"/></w:rPr>')
    assert _xml_has_marking(element) is False


def test_xml_has_marking_detects_shd_fill() -> None:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    element = parse_xml(f'<w:rPr {nsdecls("w")}><w:shd w:fill="FFFF00"/></w:rPr>')
    assert _xml_has_marking(element) is True


def test_xml_has_marking_ignores_shd_auto_fill() -> None:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    element = parse_xml(f'<w:rPr {nsdecls("w")}><w:shd w:fill="auto"/></w:rPr>')
    assert _xml_has_marking(element) is False


# ---------------------------------------------------------------------------
# extract_marked_blocks_from_docx
# ---------------------------------------------------------------------------


def test_extract_marked_blocks_from_docx_selects_highlighted_and_skips_bold(
    tmp_path: Path,
) -> None:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX

    doc = Document()
    bold_para = doc.add_paragraph()
    bold_run = bold_para.add_run("Заголовок ведущего")
    bold_run.bold = True

    doc.add_paragraph("Обычный текст без разметки")

    marked_para = doc.add_paragraph()
    marked_run = marked_para.add_run("Закадровый текст с выделением")
    marked_run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    path = tmp_path / "scenario.docx"
    doc.save(str(path))

    blocks = extract_marked_blocks_from_docx(path)

    assert blocks == ["Закадровый текст с выделением"]


def test_extract_marked_blocks_from_docx_falls_back_to_full_paragraph_for_short_fragment(
    tmp_path: Path,
) -> None:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX

    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Первая часть предложения без разметки, ")
    short_run = para.add_run("да")
    short_run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    para.add_run(" и продолжение.")

    path = tmp_path / "scenario.docx"
    doc.save(str(path))

    blocks = extract_marked_blocks_from_docx(path)

    assert blocks == [para.text.strip()]


def test_extract_marked_blocks_from_docx_uses_paragraph_level_shading(
    tmp_path: Path,
) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    para = doc.add_paragraph("Абзац выделен целиком заливкой")
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "FFFF00")
    para._p.get_or_add_pPr().append(shd)

    path = tmp_path / "scenario.docx"
    doc.save(str(path))

    blocks = extract_marked_blocks_from_docx(path)

    assert blocks == ["Абзац выделен целиком заливкой"]


# ---------------------------------------------------------------------------
# extract_paragraphs_from_docx
# ---------------------------------------------------------------------------


def test_extract_paragraphs_from_docx_reports_formatting_metadata(
    tmp_path: Path,
) -> None:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX

    doc = Document()
    doc.add_paragraph("")  # empty paragraph must be skipped entirely

    bold_para = doc.add_paragraph()
    bold_run = bold_para.add_run("Ведущий в кадре")
    bold_run.bold = True

    marked_para = doc.add_paragraph()
    marked_run = marked_para.add_run("Закадровый текст")
    marked_run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    doc.add_paragraph("Просто текст")

    path = tmp_path / "scenario.docx"
    doc.save(str(path))

    paragraphs = extract_paragraphs_from_docx(path)

    assert [p.text for p in paragraphs] == [
        "Ведущий в кадре",
        "Закадровый текст",
        "Просто текст",
    ]
    assert (paragraphs[0].is_bold, paragraphs[0].has_highlight) == (True, False)
    assert (paragraphs[1].is_bold, paragraphs[1].has_highlight) == (False, True)
    assert (paragraphs[2].is_bold, paragraphs[2].has_highlight) == (False, False)
    assert [p.index for p in paragraphs] == [1, 2, 3]


# ---------------------------------------------------------------------------
# extract_text_from_doc fallback chain
# ---------------------------------------------------------------------------


def test_doc_extraction_raises_when_no_converter_available(tmp_path: Path) -> None:
    source = tmp_path / "scenario.doc"
    source.write_bytes(b"legacy")

    with patch("stage1.io_utils.subprocess.run", side_effect=FileNotFoundError()):
        with pytest.raises(RuntimeError, match="Neither catdoc nor textutil found"):
            extract_text_from_doc(source)


def test_doc_extraction_falls_back_through_catdoc_error_to_textutil(
    tmp_path: Path,
) -> None:
    source = tmp_path / "scenario.doc"
    source.write_bytes(b"legacy")
    calls: list[str] = []

    def fake_run(cmd, **_kwargs):
        calls.append(cmd[0])
        if cmd[0] == "antiword":
            raise FileNotFoundError()
        if cmd[0] == "catdoc":
            raise subprocess.CalledProcessError(1, cmd, output="", stderr="catdoc failed")
        return CompletedProcess(cmd, 0, stdout="ЗАКАДР: Текст.\n", stderr="")

    with patch("stage1.io_utils.subprocess.run", side_effect=fake_run):
        result = extract_text_from_doc(source)

    assert calls == ["antiword", "catdoc", "textutil"]
    assert result == "ЗАКАДР: Текст."


def test_doc_extraction_raises_with_last_tool_error_when_all_fail(
    tmp_path: Path,
) -> None:
    source = tmp_path / "scenario.doc"
    source.write_bytes(b"legacy")

    def fake_run(cmd, **_kwargs):
        raise subprocess.CalledProcessError(1, cmd, output="", stderr=f"{cmd[0]} exploded")

    with patch("stage1.io_utils.subprocess.run", side_effect=fake_run):
        with pytest.raises(RuntimeError, match="textutil exploded"):
            extract_text_from_doc(source)


def test_doc_extraction_raises_with_timeout_message(tmp_path: Path) -> None:
    source = tmp_path / "scenario.doc"
    source.write_bytes(b"legacy")

    def fake_run(cmd, **_kwargs):
        raise subprocess.TimeoutExpired(cmd, 60)

    with patch("stage1.io_utils.subprocess.run", side_effect=fake_run):
        with pytest.raises(RuntimeError, match="textutil timed out after 60 seconds"):
            extract_text_from_doc(source)


# ---------------------------------------------------------------------------
# extract_text
# ---------------------------------------------------------------------------


def test_extract_text_normalizes_plain_text_and_markdown(tmp_path: Path) -> None:
    txt = tmp_path / "scenario.txt"
    txt.write_text("Привет   мир\r\n\r\n\r\nПродолжение", encoding="utf-8")
    assert extract_text(txt) == "Привет мир\n\nПродолжение"

    md = tmp_path / "scenario.md"
    md.write_text("# Заголовок", encoding="utf-8")
    assert extract_text(md) == "# Заголовок"


def test_extract_text_rejects_unsupported_format(tmp_path: Path) -> None:
    unsupported = tmp_path / "scenario.pdf"
    unsupported.write_bytes(b"%PDF")

    with pytest.raises(ValueError, match="Unsupported input format"):
        extract_text(unsupported)


def test_extract_text_dispatches_docx_and_doc_handlers(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    docx_path = tmp_path / "scenario.docx"
    docx_path.write_bytes(b"")
    doc_path = tmp_path / "scenario.doc"
    doc_path.write_bytes(b"")

    monkeypatch.setattr(
        stage1_io_utils, "extract_text_from_docx", lambda p: f"docx:{p.name}"
    )
    monkeypatch.setattr(
        stage1_io_utils, "extract_text_from_doc", lambda p: f"doc:{p.name}"
    )

    assert extract_text(docx_path) == "docx:scenario.docx"
    assert extract_text(doc_path) == "doc:scenario.doc"
